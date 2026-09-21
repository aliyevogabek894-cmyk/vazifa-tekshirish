import io
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from django.utils import timezone
from datetime import timedelta

from core.models import Classroom, Subject, StudentProfile, Assignment, StudentAssignment
from core.utils.excel_exporter import generate_homework_excel_report
from core.utils.excel_importer import import_students_from_excel


class ExcelExporterTests(TestCase):
    def setUp(self):
        self.classroom = Classroom.objects.create(name="1-A", grade_level=1)
        self.subject = Subject.objects.create(name="Matematika", code="MATH")
        self.admin = User.objects.create_superuser(username="admin", password="password123")

        self.user_st = User.objects.create_user(username="std_1", first_name="Ali", last_name="Valiyev")
        self.student = StudentProfile.objects.create(
            user=self.user_st, first_name="Ali", last_name="Valiyev",
            classroom=self.classroom, phone_number="+998901234567"
        )

        self.assignment = Assignment.objects.create(
            title="Qo'shish amali",
            subject=self.subject,
            due_date=timezone.now() + timedelta(days=2),
            created_by=self.admin
        )
        self.assignment.classrooms.add(self.classroom)

        self.task = StudentAssignment.objects.create(
            assignment=self.assignment,
            student=self.student,
            status='completed',
            marked_done_at=timezone.now()
        )

    def test_excel_export_has_all_six_sheets(self):
        """Validates that generate_homework_excel_report produces an Excel workbook with all 6 required sheets"""
        buf, filename = generate_homework_excel_report(report_type="all")
        self.assertTrue(filename.endswith(".xlsx"))

        wb = openpyxl.load_workbook(buf)
        sheet_names = wb.sheetnames

        expected_sheets = [
            "Umumiy hisobot",
            "O'quvchilar statistikasi",
            "Vazifalar batafsil",
            "Bajarilgan vazifalar",
            "Bajarilmagan vazifalar",
            "Fayllar va topshiriqlar"
        ]

        for expected in expected_sheets:
            self.assertIn(expected, sheet_names, f"{expected} varag'i topilmadi!")

        # Verify sheet 2 has the student
        ws2 = wb["O'quvchilar statistikasi"]
        self.assertGreaterEqual(ws2.max_row, 2)
        row2_values = [cell.value for cell in ws2[2]]
        self.assertIn("Ali Valiyev", row2_values)

    def test_excel_download_view_authenticated(self):
        """Admin can download the report via the view"""
        client = Client()
        client.force_login(self.admin)

        url = reverse('admin_export_excel')
        response = client.get(url, {'report_type': 'single_class', 'classroom_id': self.classroom.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment;', response['Content-Disposition'])

    def test_bulk_excel_import(self):
        """Tests importing students from an in-memory Excel workbook"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Ism", "Familiya", "Sinf", "Telefon"])
        ws.append(["Jasur", "Olimov", "1-A", "+998909876543"])
        ws.append(["Dilnoza", "Karimova", "1-A", "+998908765432"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        result = import_students_from_excel(buf)
        self.assertEqual(result['success_count'], 2)
        self.assertEqual(len(result['errors']), 0)

        # Check in DB
        self.assertTrue(StudentProfile.objects.filter(phone_number="+998909876543").exists())
        self.assertTrue(StudentProfile.objects.filter(phone_number="+998908765432").exists())
