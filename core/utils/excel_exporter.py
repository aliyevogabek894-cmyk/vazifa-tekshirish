import io
from datetime import datetime, date
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import StudentAssignment, StudentProfile, Classroom, Assignment, SubmissionAttachment


def style_header_cell(cell, text):
    cell.value = text
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def get_status_style(status):
    """Returns (fill, font_color, display_text) for a status"""
    status_map = {
        'approved': ("DCFCE7", "166534", "Qabul qilindi"),
        'completed': ("DCFCE7", "166534", "Bajarildi"),
        'submitted': ("DBEAFE", "1E40AF", "Javob yuborildi"),
        'under_review': ("FEF3C7", "92400E", "Tekshirilmoqda"),
        'needs_work': ("FFEDD5", "9A3412", "Qayta ishlash kerak"),
        'overdue': ("FEE2E2", "991B1B", "Muddati o‘tgan"),
        'not_started': ("F3F4F6", "374151", "Bajarilmadi"),
    }
    return status_map.get(status, ("FFFFFF", "000000", status))


def autofit_columns(ws, max_cols=30):
    for col in ws.iter_cols(min_col=1, max_col=min(ws.max_column, max_cols)):
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and '%' in cell.number_format:
                val_str += ' %'
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_homework_excel_report(
    report_type: str = "all", # 'single_student', 'single_class', 'all'
    classroom_id: int = None,
    student_id: int = None,
    subject_id: int = None,
    start_date: str = None,
    end_date: str = None,
    status_filter: str = None
) -> tuple[io.BytesIO, str]:
    """
    Generates a comprehensive 6-sheet professional Excel report using openpyxl.
    Returns: (BytesIO buffer, filename)
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Base Queryset
    qs = StudentAssignment.objects.select_related(
        'student', 'student__classroom', 'assignment', 'assignment__subject', 'submission'
    ).prefetch_related('submission__attachments')

    if classroom_id:
        qs = qs.filter(student__classroom_id=classroom_id)
    if student_id:
        qs = qs.filter(student_id=student_id)
    if subject_id:
        qs = qs.filter(assignment__subject_id=subject_id)
    if start_date:
        qs = qs.filter(assignment__assigned_date__gte=start_date)
    if end_date:
        qs = qs.filter(assignment__assigned_date__lte=end_date)
    if status_filter and status_filter != 'all':
        qs = qs.filter(status=status_filter)

    records = list(qs.order_by('student__classroom__name', 'student__last_name', '-assignment__due_date'))

    # Thin border helper
    thin = Side(border_style="thin", color="CBD5E1")
    border_all = Border(top=thin, left=thin, right=thin, bottom=thin)

    # Determine report name and filename
    date_range_str = f"{start_date or 'boshidan'}_{end_date or 'hozirgacha'}"
    if report_type == 'single_student' and student_id:
        student = StudentProfile.objects.filter(id=student_id).first()
        name_part = f"{student.last_name}_{student.first_name}" if student else f"oquvchi_{student_id}"
        filename = f"{name_part}_hisobot_{date_range_str}.xlsx"
        report_title = f"O'quvchi hisoboti: {student.full_name if student else ''}"
    elif report_type == 'single_class' and classroom_id:
        classroom = Classroom.objects.filter(id=classroom_id).first()
        class_name = classroom.name if classroom else f"sinf_{classroom_id}"
        filename = f"{class_name}_hisobot_{date_range_str}.xlsx"
        report_title = f"{class_name} sinfi hisoboti"
    else:
        filename = f"Barcha_sinflar_hisobot_{date_range_str}.xlsx"
        report_title = "Barcha sinflar bo'yicha umumiy hisobot"

    # -------------------------------------------------------------
    # SHEET 1: Umumiy hisobot
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Umumiy hisobot")
    ws1.views.sheetView[0].showGridLines = True

    # Title block
    ws1.merge_cells("A1:F1")
    title_cell = ws1["A1"]
    title_cell.value = "MAKTAB UY VAZIFALARI MONITORING HISOBOTI"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    ws1.merge_cells("A2:F2")
    ws1["A2"].value = f"{report_title} | Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws1["A2"].font = Font(name="Calibri", size=11, italic=True, color="4B5563")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # Summary Metrics Calculation
    total_assigned = len(records)
    completed_count = sum(1 for r in records if r.status in ['completed', 'approved'])
    submitted_count = sum(1 for r in records if r.status == 'submitted')
    under_review_count = sum(1 for r in records if r.status == 'under_review')
    approved_count = sum(1 for r in records if r.status == 'approved')
    needs_work_count = sum(1 for r in records if r.status == 'needs_work')
    not_started_count = sum(1 for r in records if r.status in ['not_started', 'overdue'])

    distinct_students = {r.student_id for r in records}
    distinct_classes = {r.student.classroom_id for r in records if r.student.classroom_id}
    pct_done = (completed_count / total_assigned * 100) if total_assigned > 0 else 0

    metrics = [
        ("Hisobot davri:", f"{start_date or 'Boshlang\'ich sana'} — {end_date or 'Hozirgi sana'}"),
        ("Jami jalb qilingan sinflar soni:", len(distinct_classes)),
        ("Jami o'quvchilar soni:", len(distinct_students)),
        ("Jami berilgan topshiriqlar (o'quvchi-vazifa):", total_assigned),
        ("Bajarilgan va qabul qilingan vazifalar:", completed_count),
        ("Tekshiruv kutilayotgan javoblar:", submitted_count + under_review_count),
        ("Qayta ishlashga qaytarilgan:", needs_work_count),
        ("Bajarilmagan / Muddati o'tgan vazifalar:", not_started_count),
        ("Umumiy bajarish foizi:", f"{pct_done:.1f}%"),
    ]

    ws1.append([]) # row 3 blank
    ws1.append(["Ko'rsatkich", "Qiymat"])
    style_header_cell(ws1.cell(row=4, column=1), "Asosiy ko'rsatkich")
    style_header_cell(ws1.cell(row=4, column=2), "Miqdor / Holat")
    ws1.row_dimensions[4].height = 25

    row_idx = 5
    for label, val in metrics:
        c1 = ws1.cell(row=row_idx, column=1, value=label)
        c2 = ws1.cell(row=row_idx, column=2, value=val)
        c1.font = Font(name="Calibri", size=11, bold=(row_idx in [5, 8, 13]))
        c2.font = Font(name="Calibri", size=11, bold=True)
        c1.border = border_all
        c2.border = border_all
        if label == "Umumiy bajarish foizi:":
            c2.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            c2.font = Font(name="Calibri", size=12, bold=True, color="166534")
        row_idx += 1

    autofit_columns(ws1)

    # -------------------------------------------------------------
    # SHEET 2: O‘quvchilar statistikasi
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="O'quvchilar statistikasi")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "A2"

    headers2 = [
        "№", "F.I.O", "Sinf", "Telefon raqami", "Jami vazifalar",
        "Bajarildi", "Bajarilmadi", "Javob yuborildi", "Qabul qilindi",
        "Qayta ishlash", "Bajarish foizi"
    ]
    ws2.append(headers2)
    for col_i, h in enumerate(headers2, start=1):
        style_header_cell(ws2.cell(row=1, column=col_i), h)
    ws2.row_dimensions[1].height = 28

    # Group records by student
    student_records = {}
    for r in records:
        sid = r.student_id
        if sid not in student_records:
            student_records[sid] = {
                'student': r.student,
                'total': 0,
                'completed': 0,
                'not_started': 0,
                'submitted': 0,
                'approved': 0,
                'needs_work': 0,
            }
        data = student_records[sid]
        data['total'] += 1
        if r.status in ['completed', 'approved']:
            data['completed'] += 1
        if r.status in ['not_started', 'overdue']:
            data['not_started'] += 1
        if r.status in ['submitted', 'under_review']:
            data['submitted'] += 1
        if r.status == 'approved':
            data['approved'] += 1
        if r.status == 'needs_work':
            data['needs_work'] += 1

    row_num = 2
    for idx, (sid, st_data) in enumerate(student_records.items(), start=1):
        st = st_data['student']
        tot = st_data['total']
        done = st_data['completed']
        pct = (done / tot) if tot > 0 else 0.0

        row_vals = [
            idx,
            st.full_name,
            st.classroom.name if st.classroom else "-",
            st.phone_number,
            tot,
            done,
            st_data['not_started'],
            st_data['submitted'],
            st_data['approved'],
            st_data['needs_work'],
            f"{pct * 100:.1f}%"
        ]
        ws2.append(row_vals)
        for col_i in range(1, len(headers2) + 1):
            c = ws2.cell(row=row_num, column=col_i)
            c.border = border_all
            c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 3, 5, 6, 7, 8, 9, 10, 11] else "left")
            if col_i == 11:
                # Color code percentage
                if pct >= 0.8:
                    c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    c.font = Font(name="Calibri", bold=True, color="166534")
                elif pct < 0.5:
                    c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    c.font = Font(name="Calibri", bold=True, color="991B1B")
                else:
                    c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    c.font = Font(name="Calibri", bold=True, color="92400E")
        row_num += 1

    ws2.auto_filter.ref = ws2.dimensions
    autofit_columns(ws2)

    # -------------------------------------------------------------
    # SHEET 3: Vazifalar bo‘yicha batafsil hisobot
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Vazifalar batafsil")
    ws3.views.sheetView[0].showGridLines = True
    ws3.freeze_panes = "A2"

    headers3 = [
        "№", "O'quvchi", "Sinf", "Fan", "Vazifa nomi", "Berilgan sana",
        "Muddati", "Holati", "Bajarilgan vaqt", "O'qituvchi xulosasi", "O'qituvchi izohi"
    ]
    ws3.append(headers3)
    for col_i, h in enumerate(headers3, start=1):
        style_header_cell(ws3.cell(row=1, column=col_i), h)
    ws3.row_dimensions[1].height = 28

    row_num = 2
    for idx, r in enumerate(records, start=1):
        submission = getattr(r, 'submission', None)
        teacher_feedback = submission.teacher_feedback if submission else ""
        review_stat = submission.get_review_status_display() if submission else "-"
        marked_time = r.marked_done_at.strftime("%d.%m.%Y %H:%M") if r.marked_done_at else "-"

        fill_col, font_col, status_text = get_status_style(r.current_status)

        row_vals = [
            idx,
            r.student.full_name,
            r.student.classroom.name if r.student.classroom else "-",
            r.assignment.subject.name,
            r.assignment.title,
            r.assignment.assigned_date.strftime("%d.%m.%Y") if r.assignment.assigned_date else "-",
            r.assignment.due_date.strftime("%d.%m.%Y %H:%M") if r.assignment.due_date else "-",
            status_text,
            marked_time,
            review_stat,
            teacher_feedback
        ]
        ws3.append(row_vals)
        for col_i in range(1, len(headers3) + 1):
            c = ws3.cell(row=row_num, column=col_i)
            c.border = border_all
            c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 3, 6, 7, 8, 9, 10] else "left")
            if col_i == 8:
                c.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
                c.font = Font(name="Calibri", bold=True, color=font_col)
        row_num += 1

    ws3.auto_filter.ref = ws3.dimensions
    autofit_columns(ws3)

    # -------------------------------------------------------------
    # SHEET 4: Bajarilgan vazifalar
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Bajarilgan vazifalar")
    ws4.views.sheetView[0].showGridLines = True
    ws4.freeze_panes = "A2"

    headers4 = ["№", "O'quvchi", "Sinf", "Fan", "Vazifa nomi", "Bajarilgan sana va vaqt", "Holati", "Fayllar soni"]
    ws4.append(headers4)
    for col_i, h in enumerate(headers4, start=1):
        style_header_cell(ws4.cell(row=1, column=col_i), h)
    ws4.row_dimensions[1].height = 28

    completed_records = [r for r in records if r.status in ['completed', 'approved', 'submitted', 'under_review']]
    row_num = 2
    for idx, r in enumerate(completed_records, start=1):
        submission = getattr(r, 'submission', None)
        file_count = submission.attachments.count() if submission else 0
        marked_time = r.marked_done_at.strftime("%d.%m.%Y %H:%M") if r.marked_done_at else "-"
        fill_col, font_col, status_text = get_status_style(r.current_status)

        ws4.append([
            idx,
            r.student.full_name,
            r.student.classroom.name if r.student.classroom else "-",
            r.assignment.subject.name,
            r.assignment.title,
            marked_time,
            status_text,
            file_count
        ])
        for col_i in range(1, len(headers4) + 1):
            c = ws4.cell(row=row_num, column=col_i)
            c.border = border_all
            c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 3, 6, 7, 8] else "left")
            if col_i == 7:
                c.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
                c.font = Font(name="Calibri", bold=True, color=font_col)
        row_num += 1

    ws4.auto_filter.ref = ws4.dimensions
    autofit_columns(ws4)

    # -------------------------------------------------------------
    # SHEET 5: Bajarilmagan vazifalar
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Bajarilmagan vazifalar")
    ws5.views.sheetView[0].showGridLines = True
    ws5.freeze_panes = "A2"

    headers5 = ["№", "O'quvchi", "Sinf", "Telefon raqami", "Fan", "Vazifa nomi", "Topshirish muddati", "Holat"]
    ws5.append(headers5)
    for col_i, h in enumerate(headers5, start=1):
        style_header_cell(ws5.cell(row=1, column=col_i), h)
    ws5.row_dimensions[1].height = 28

    uncompleted_records = [r for r in records if r.status in ['not_started', 'overdue', 'needs_work']]
    row_num = 2
    for idx, r in enumerate(uncompleted_records, start=1):
        fill_col, font_col, status_text = get_status_style(r.current_status)
        ws5.append([
            idx,
            r.student.full_name,
            r.student.classroom.name if r.student.classroom else "-",
            r.student.phone_number,
            r.assignment.subject.name,
            r.assignment.title,
            r.assignment.due_date.strftime("%d.%m.%Y %H:%M") if r.assignment.due_date else "-",
            status_text
        ])
        for col_i in range(1, len(headers5) + 1):
            c = ws5.cell(row=row_num, column=col_i)
            c.border = border_all
            c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 3, 4, 7, 8] else "left")
            if col_i == 8:
                c.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
                c.font = Font(name="Calibri", bold=True, color=font_col)
        row_num += 1

    ws5.auto_filter.ref = ws5.dimensions
    autofit_columns(ws5)

    # -------------------------------------------------------------
    # SHEET 6: Yuborilgan fayllar va topshiriqlar
    # -------------------------------------------------------------
    ws6 = wb.create_sheet(title="Fayllar va topshiriqlar")
    ws6.views.sheetView[0].showGridLines = True
    ws6.freeze_panes = "A2"

    headers6 = [
        "№", "O'quvchi", "Sinf", "Vazifa", "Fayl nomi", "Fayl turi",
        "Fayl hajmi", "Yuklangan vaqt", "O'qituvchi izohi"
    ]
    ws6.append(headers6)
    for col_i, h in enumerate(headers6, start=1):
        style_header_cell(ws6.cell(row=1, column=col_i), h)
    ws6.row_dimensions[1].height = 28

    row_num = 2
    file_idx = 1
    for r in records:
        submission = getattr(r, 'submission', None)
        if submission and submission.attachments.exists():
            for att in submission.attachments.all():
                ws6.append([
                    file_idx,
                    r.student.full_name,
                    r.student.classroom.name if r.student.classroom else "-",
                    r.assignment.title,
                    att.file_name,
                    att.get_file_type_display(),
                    att.size_display,
                    att.uploaded_at.strftime("%d.%m.%Y %H:%M") if att.uploaded_at else "-",
                    submission.teacher_feedback or "-"
                ])
                for col_i in range(1, len(headers6) + 1):
                    c = ws6.cell(row=row_num, column=col_i)
                    c.border = border_all
                    c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 3, 6, 7, 8] else "left")
                file_idx += 1
                row_num += 1

    ws6.auto_filter.ref = ws6.dimensions
    autofit_columns(ws6)

    # Save to BytesIO stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, filename
