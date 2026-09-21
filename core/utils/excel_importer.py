import openpyxl
import re
from django.contrib.auth.models import User
from django.db import transaction
from core.models import StudentProfile, Classroom
from core.utils.otp_service import normalize_phone


def import_students_from_excel(file_obj) -> dict:
    """
    Imports students from an uploaded Excel file.
    Expected Columns in Row 1:
    - Ism (First Name)
    - Familiya (Last Name)
    - Sinf (Classroom Name, e.g. '1-A')
    - Telefon (Phone number, e.g. '+998901234567')

    Returns: {'success_count': int, 'updated_count': int, 'errors': list[str]}
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    headers = [str(cell.value or '').strip().lower() for cell in ws[1]]
    
    # Map header names
    first_name_col = None
    last_name_col = None
    class_col = None
    phone_col = None

    for idx, h in enumerate(headers):
        if 'ism' in h or 'first' in h:
            first_name_col = idx
        elif 'familiya' in h or 'last' in h:
            last_name_col = idx
        elif 'sinf' in h or 'class' in h:
            class_col = idx
        elif 'telefon' in h or 'phone' in h or 'raqam' in h:
            phone_col = idx

    if phone_col is None:
        return {'success_count': 0, 'updated_count': 0, 'errors': ["Excel faylida 'Telefon' ustuni topilmadi!"]}

    success_count = 0
    updated_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue # empty row

        first_name = str(row[first_name_col] or '').strip() if first_name_col is not None else "O'quvchi"
        last_name = str(row[last_name_col] or '').strip() if last_name_col is not None else ""
        class_name = str(row[class_col] or '').strip() if class_col is not None else ""
        raw_phone = str(row[phone_col] or '').strip()

        if not raw_phone:
            errors.append(f"{row_idx}-qator: Telefon raqami ko'rsatilmagan.")
            continue

        clean_phone = normalize_phone(raw_phone)
        if len(re.sub(r'\D', '', clean_phone)) < 9:
            errors.append(f"{row_idx}-qator: Telefon raqami noto'g'ri: {raw_phone}")
            continue

        try:
            with transaction.atomic():
                # Find or create classroom
                classroom = None
                if class_name:
                    classroom, _ = Classroom.objects.get_or_create(
                        name=class_name,
                        defaults={'grade_level': int(re.search(r'\d+', class_name).group(0)) if re.search(r'\d+', class_name) else 1}
                    )

                # Check if StudentProfile with this phone already exists
                student = StudentProfile.objects.filter(phone_number=clean_phone).first()
                if student:
                    student.first_name = first_name or student.first_name
                    student.last_name = last_name or student.last_name
                    if classroom:
                        student.classroom = classroom
                    student.save()
                    updated_count += 1
                else:
                    # Create User
                    username = f"std_{re.sub(r'[^0-9]', '', clean_phone)[-9:]}"
                    # Make username unique if collision
                    base_username = username
                    counter = 1
                    while User.objects.filter(username=username).exists():
                        username = f"{base_username}_{counter}"
                        counter += 1

                    user = User.objects.create_user(
                        username=username,
                        first_name=first_name,
                        last_name=last_name
                    )
                    user.set_unusable_password()
                    user.save()

                    StudentProfile.objects.create(
                        user=user,
                        first_name=first_name,
                        last_name=last_name,
                        classroom=classroom,
                        phone_number=clean_phone
                    )
                    success_count += 1
        except Exception as e:
            errors.append(f"{row_idx}-qator: Xatolik - {str(e)}")

    return {'success_count': success_count, 'updated_count': updated_count, 'errors': errors}
